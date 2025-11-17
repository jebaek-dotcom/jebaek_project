"""
DOCX 제안서 생성 모듈 (고도화 버전)
기업 분석 및 마케팅 전략 보고서
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from typing import Dict, List
import os


class ProposalGenerator:
    def __init__(self, analysis_data: Dict):
        self.data = analysis_data
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """문서 스타일 설정"""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = '맑은 고딕'
        font.size = Pt(11)

    def generate(self, output_path: str) -> str:
        """보고서 생성"""
        # 표지
        self._add_cover_page()
        self.doc.add_page_break()

        # 목차
        self._add_table_of_contents()
        self.doc.add_page_break()

        # 1. 핵심 요약
        self._add_executive_summary()
        self.doc.add_page_break()

        # 2. 기업 개요
        self._add_company_overview()
        self.doc.add_page_break()

        # 3. 핵심 사업 영역
        self._add_core_business()
        self.doc.add_page_break()

        # 4. 타겟 시장 & 고객
        self._add_target_market()
        self.doc.add_page_break()

        # 5. 경쟁사 분석
        self._add_competitor_analysis()
        self.doc.add_page_break()

        # 6. 시장 및 고객 행동 분석
        self._add_market_behavior()
        self.doc.add_page_break()

        # 7. 산업 가치 + 기존 마케팅 한계
        self._add_industry_limits()
        self.doc.add_page_break()

        # 8. 온드미디어 전략
        self._add_owned_media_strategy()
        self.doc.add_page_break()

        # 9. Plette Agent 제안
        self._add_plette_agent_proposal()
        self.doc.add_page_break()

        # 10. 기대 효과 및 결론
        self._add_expected_outcomes()
        self.doc.add_page_break()

        # 11. Next Steps & 로드맵
        self._add_implementation_roadmap()

        # 저장
        self.doc.save(output_path)
        return output_path

    def _add_cover_page(self):
        """표지"""
        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(f'\n\n\n\n{self.data.get("document_title", "기업 분석 및 마케팅 전략 보고서")}\n\n')
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f'{self.data.get("company_name", "기업명")}\n')
        run.font.size = Pt(20)
        run.font.bold = True

        date_para = self.doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = date_para.add_run(f'\n\n\n{datetime.now().strftime("%Y년 %m월 %d일")}')
        run.font.size = Pt(14)

    def _add_table_of_contents(self):
        """목차"""
        heading = self.doc.add_heading('목차', level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

        toc_items = [
            '1. 핵심 요약 (Executive Summary)',
            '2. 기업 개요',
            '3. 핵심 사업 영역',
            '4. 현재 타겟 시장 & 고객',
            '5. 경쟁사 분석',
            '6. 시장 및 고객 행동 분석',
            '7. 산업·업계 가치 + 기존 마케팅의 한계',
            '8. 온드미디어 중심 마케팅 전략 제안',
            '9. Plette Agent 활용 제안',
            '10. 기대 효과 및 결론',
            '11. Next Step 제안 및 도입 로드맵'
        ]

        for item in toc_items:
            p = self.doc.add_paragraph(item, style='List Number')
            p.paragraph_format.left_indent = Inches(0.5)

    def _add_executive_summary(self):
        """1. 핵심 요약"""
        self.doc.add_heading('1. 핵심 요약 (Executive Summary)', level=1)

        summary = self.data.get('executive_summary', {})

        p = self.doc.add_paragraph()
        p.add_run('⚡ 의사결정권자를 위한 핵심 정리\n\n').bold = True

        for key, value in summary.items():
            if key == '핵심 발견사항' and isinstance(value, list):
                p = self.doc.add_paragraph()
                p.add_run(f'{key}:\n').bold = True
                for item in value:
                    self.doc.add_paragraph(f'• {item}', style='List Bullet')
            else:
                p = self.doc.add_paragraph()
                p.add_run(f'{key}: ').bold = True
                p.add_run(str(value))

    def _add_company_overview(self):
        """2. 기업 개요"""
        self.doc.add_heading('2. 기업 개요', level=1)

        overview = self.data.get('company_overview', {})

        for key, value in overview.items():
            p = self.doc.add_paragraph()
            p.add_run(f'{key}: ').bold = True
            p.add_run(str(value))

    def _add_core_business(self):
        """3. 핵심 사업 영역"""
        self.doc.add_heading('3. 핵심 사업 영역', level=1)

        core_business = self.data.get('core_business', {})

        for key, value in core_business.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                for item in value:
                    self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}: ').bold = True
                    p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

    def _add_target_market(self):
        """4. 타겟 시장 & 고객"""
        self.doc.add_heading('4. 현재 타겟 시장 & 고객', level=1)

        target_market = self.data.get('target_market', {})

        for key, value in target_market.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    # 구조화된 데이터 (세분화, 페르소나 등)
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()  # 간격
                else:
                    # 단순 리스트
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            else:
                self.doc.add_paragraph(str(value))

    def _add_competitor_analysis(self):
        """5. 경쟁사 분석"""
        self.doc.add_heading('5. 경쟁사 분석', level=1)

        competitor = self.data.get('competitor_analysis', {})

        for key, value in competitor.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    # 구조화된 경쟁사 마케팅 데이터
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            else:
                self.doc.add_paragraph(str(value))

    def _add_market_behavior(self):
        """6. 시장 및 고객 행동 분석"""
        self.doc.add_heading('6. 시장 및 고객 행동 분석', level=1)

        behavior = self.data.get('market_customer_behavior', {})

        for key, value in behavior.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}: ').bold = True
                    if isinstance(v, list):
                        self.doc.add_paragraph()
                        for sub_item in v:
                            self.doc.add_paragraph(f'  • {sub_item}', style='List Bullet')
                    else:
                        p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

    def _add_industry_limits(self):
        """7. 산업 가치 + 기존 마케팅 한계"""
        self.doc.add_heading('7. 산업·업계 가치 + 기존 마케팅의 한계', level=1)

        p = self.doc.add_paragraph()
        p.add_run('🔥 왜 새로운 전략과 Agent가 필요한가\n').bold = True
        p.add_run('(기업의 문제 정의 파트 - 매우 중요)\n\n').italic = True

        limits = self.data.get('industry_value_marketing_limits', {})

        for key, value in limits.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            else:
                self.doc.add_paragraph(str(value))

    def _add_owned_media_strategy(self):
        """8. 온드미디어 전략"""
        self.doc.add_heading('8. 온드미디어 중심 마케팅 전략 제안', level=1)

        strategy = self.data.get('owned_media_strategy', {})

        for key, value in strategy.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    # B2B Funnel 등 구조화된 전략
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                # SNS 전략, 블로그 전략 등
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}: ').bold = True
                    if isinstance(v, list):
                        self.doc.add_paragraph()
                        for sub_item in v:
                            self.doc.add_paragraph(f'  • {sub_item}', style='List Bullet')
                    else:
                        p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

    def _add_plette_agent_proposal(self):
        """9. Plette Agent 제안"""
        self.doc.add_heading('9. Plette Agent 활용 제안', level=1)

        p = self.doc.add_paragraph()
        p.add_run('🤖 AI Agent 기반 마케팅 자동화 솔루션\n\n').bold = True

        agent = self.data.get('plette_agent_proposal', {})

        for key, value in agent.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                # 비용 절감 효과, Agent 구조 등
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}:\n').bold = True

                    if isinstance(v, dict):
                        for sub_k, sub_v in v.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'  {sub_k}: ').bold = True
                            p.add_run(str(sub_v))
                    elif isinstance(v, list):
                        for sub_item in v:
                            self.doc.add_paragraph(f'  • {sub_item}', style='List Bullet')
                    else:
                        p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

    def _add_expected_outcomes(self):
        """10. 기대 효과 및 결론"""
        self.doc.add_heading('10. 기대 효과 및 결론', level=1)

        outcomes = self.data.get('expected_outcomes', {})

        for key, value in outcomes.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                for item in value:
                    self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}: ').bold = True
                    p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

    def _add_implementation_roadmap(self):
        """11. Next Steps & 로드맵"""
        self.doc.add_heading('11. Next Step 제안 및 도입 로드맵', level=1)

        p = self.doc.add_paragraph()
        p.add_run('🚀 실행 로드맵 및 투자 계획\n\n').bold = True

        roadmap = self.data.get('next_steps_roadmap', {})

        for key, value in roadmap.items():
            self.doc.add_heading(key, level=2)

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    for item in value:
                        for k, v in item.items():
                            p = self.doc.add_paragraph()
                            p.add_run(f'{k}: ').bold = True
                            p.add_run(str(v))
                        self.doc.add_paragraph()
                else:
                    for item in value:
                        self.doc.add_paragraph(f'• {item}', style='List Bullet')
            elif isinstance(value, dict):
                for k, v in value.items():
                    p = self.doc.add_paragraph()
                    p.add_run(f'{k}: ').bold = True

                    if isinstance(v, list):
                        self.doc.add_paragraph()
                        for sub_item in v:
                            self.doc.add_paragraph(f'  • {sub_item}', style='List Bullet')
                    else:
                        p.add_run(str(v))
            else:
                self.doc.add_paragraph(str(value))

        # 마지막 감사 인사
        self.doc.add_paragraph('\n')
        thanks = self.doc.add_paragraph('본 보고서를 검토해 주셔서 감사합니다.\n상세한 논의를 위해 언제든 연락 주시기 바랍니다.')
        thanks.alignment = WD_ALIGN_PARAGRAPH.RIGHT
